"""Structured text extraction. OCR models are local and loaded only when needed.

`parse_structured()` returns units (pages / slides / paragraphs / line blocks)
with heading marks, extracted figures and a plain-text rendition; see
structure.py and docs/API.md (sidecar protocol) for the format.
"""
from contextlib import closing
from dataclasses import replace
from io import BytesIO
from pathlib import Path
import json
import re
import sys

import chapters
from figures import Figures
from structure import Block, Unit, unicode_issues


class DocumentError(ValueError):
    """A problem with the user's file that should be shown as-is."""


class Ocr:
    def __init__(self, mode="auto", cache_dir=None):
        self.engine = None
        self.mode = mode
        self.active_mode = None
        self.confidence = 1.0
        self.cache_dir = Path(cache_dir) / "images" if cache_dir else None

    def lines(self, image):
        from ocr_runtime import make_engine, provider_for
        if self.engine is None:
            self.active_mode = provider_for(self.mode)
            try:
                self.engine = make_engine(self.active_mode)
            except Exception:
                if self.active_mode == "cpu":
                    raise
                self.active_mode = "cpu"
                self.engine = make_engine("cpu")
        try:
            result = self.engine(image)
        except Exception:
            if self.active_mode == "cpu":
                raise
            self.active_mode = "cpu"
            self.engine = make_engine("cpu")
            result = self.engine(image)
        scores = getattr(result, "scores", None)
        self.confidence = sum(scores) / len(scores) if scores is not None and len(scores) else 0
        if result.txts is None:
            return []
        return list(zip(result.boxes, result.txts))

    def image_text(self, blob):
        from PIL import Image
        import hashlib
        from checkpoints import atomic_json, SCHEMA, model_fingerprint
        target = None
        if self.cache_dir:
            key = hashlib.sha256(blob + f"{SCHEMA}:{self.mode}:{model_fingerprint()}".encode()).hexdigest()
            target = self.cache_dir / f"{key}.json"
            try:
                cached = json.loads(target.read_text(encoding="utf-8"))
                if isinstance(cached, str):
                    return cached
            except (OSError, ValueError):
                pass
        with Image.open(BytesIO(blob)) as image:
            with image.convert("RGB") as rgb:
                text = "\n".join(text for _, text in self.lines(rgb))
        if target:
            atomic_json(target, text)
        return text


class Progress:
    """Reports `{"progress": {...}}` JSON lines on stderr when enabled."""

    def __init__(self, enabled=False):
        self.enabled = enabled

    def __call__(self, current, total, stage="page"):
        if self.enabled:
            print(json.dumps({"progress": {"stage": stage, "current": current, "total": total}}),
                  file=sys.stderr, flush=True)

    def page(self, unit, images, total):
        if self.enabled:
            print(json.dumps({"page_result": {"page": unit.to_json(), "images": images, "total": total}},
                             ensure_ascii=False), file=sys.stderr, flush=True)


def normalized(text):
    return re.sub(r"\W+", "", text, flags=re.UNICODE).casefold()


def same_text(recognized, existing):
    # Case, operators, superscripts and mathematical alphabets carry meaning.
    # Only whitespace is ignored, and corrupt mappings never suppress OCR.
    return (not unicode_issues(existing) and bool(existing.strip())
            and re.sub(r'\s+', '', recognized) == re.sub(r'\s+', '', existing))


# --------------------------------------------------------------------------- PDF

def _line_sizes(textpage):
    """Maps normalized line text to its font size, sampled at the start of each text rect."""
    import pypdfium2.raw as raw

    sizes = {}
    try:
        for index in range(textpage.count_rects()):
            left, bottom, right, top = textpage.get_rect(index)
            text = textpage.get_text_bounded(left, bottom, right, top).strip()
            key = normalized(text)
            if not key:
                continue
            char = textpage.get_index(left + 1, (bottom + top) / 2, 3, 3)
            size = raw.FPDFText_GetFontSize(textpage, char) if char is not None and char >= 0 else 0
            sizes[key] = max(sizes.get(key, 0.0), size or (top - bottom))
    except Exception:
        return {}
    return sizes


def _read_toc(pdf):
    toc = []
    try:
        for bookmark in pdf.get_toc(max_depth=4):
            dest = bookmark.get_dest()
            toc.append((bookmark.level, bookmark.get_title(), dest.get_index() if dest else None))
    except Exception:
        return []
    return toc


def _native_blocks(textpage):
    """Read full Unicode scalars and their boxes; never decode through UCS-2."""
    import pypdfium2.raw as raw
    blocks, chars, boxes, sizes = [], [], [], []

    def flush():
        # Some ToUnicode maps expose non-BMP characters as two UTF-16 units.
        # Join valid pairs, while retaining a visible replacement for broken ones.
        text = ''.join(chars).encode('utf-16-le', errors='surrogatepass').decode('utf-16-le', errors='replace')
        if text.strip():
            bbox = ([min(b[0] for b in boxes), min(b[1] for b in boxes),
                     max(b[2] for b in boxes), max(b[3] for b in boxes)] if boxes else None)
            blocks.append(Block(text.strip(), raw_text=text, bbox=bbox,
                                size=sorted(sizes)[len(sizes) // 2] if sizes else None))
        chars.clear()
        boxes.clear()
        sizes.clear()

    for index in range(textpage.count_chars()):
        scalar = raw.FPDFText_GetUnicode(textpage, index)
        char = chr(scalar) if 0 < scalar <= 0x10ffff else '\ufffd'
        if char in '\r\n':
            flush()
            continue
        chars.append(char)
        if not char.isspace():
            boxes.append(textpage.get_charbox(index))
            size = raw.FPDFText_GetFontSize(textpage, index)
            if size > 0:
                sizes.append(size)
    flush()
    return blocks


def parse_pdf(path, ocr, figures, progress, cache_dir=None):
    from pdf_pipeline import parse_pdf as run
    return run(path, ocr, figures, progress, cache_dir)


TERMINAL = tuple("。！？；：.!?;:」』”）)")


def _is_cjk(char):
    return "\u3040" <= char <= "\u30ff" or "\u3400" <= char <= "\u9fff" or "\uf900" <= char <= "\ufaff"


def merge_lines(unit):
    """Joins PDF lines that were wrapped by the page layout back into paragraphs."""
    lengths = sorted(len(b.text) for b in unit.blocks if b.kind == "text" and not b.heading and len(b.text) >= 10)
    typical = lengths[len(lengths) * 3 // 4] if lengths else 0
    merged = []
    for block in unit.blocks:
        previous = merged[-1] if merged else None
        joinable = (
            previous is not None and typical
            and previous.kind == "text" and block.kind == "text"
            and previous.heading is None and block.heading is None
            and not previous.formula and not block.formula
            and not previous.issues and not block.issues
            and not previous.text.endswith(TERMINAL)
            and not chapters.TOC_LINE.search(previous.text) and not chapters.TOC_LINE.search(block.text)
            and len(previous.text.splitlines()[-1]) >= typical * 0.8
            and not (previous.size and block.size and abs(previous.size - block.size) > 1.5)
        )
        if joinable:
            previous.raw_text += "\n" + block.raw_text
            if previous.bbox and block.bbox:
                previous.bbox = [min(previous.bbox[0], block.bbox[0]), min(previous.bbox[1], block.bbox[1]),
                                 max(previous.bbox[2], block.bbox[2]), max(previous.bbox[3], block.bbox[3])]
            tail, head = previous.text[-1], block.text[0]
            if tail == "-" and head.islower():
                previous.text = previous.text[:-1] + block.text
            elif _is_cjk(tail) or _is_cjk(head):
                previous.text += block.text
            else:
                previous.text += " " + block.text
        else:
            merged.append(replace(block))
    unit.blocks = merged


# --------------------------------------------------------------------------- Word

def _docx_heading_level(paragraph, qn, style_names):
    properties = paragraph.find(qn("w:pPr"))
    if properties is None:
        return None
    outline = properties.find(qn("w:outlineLvl"))
    if outline is not None:
        try:
            level = int(outline.get(qn("w:val")))
            if 0 <= level <= 5:
                return level + 1
        except (TypeError, ValueError):
            pass
    style = properties.find(qn("w:pStyle"))
    if style is None:
        return None
    name = style_names.get(style.get(qn("w:val")), "").strip().lower()
    match = re.match(r"^(?:heading|标题)\s*(\d)$", name)
    return int(match.group(1)) if match else None


def parse_docx(path, ocr, figures, progress):
    from docx import Document
    from docx.oxml.ns import qn

    try:
        document = Document(path)
    except Exception as error:
        raise DocumentError(f"无法读取 Word 文件（可能已损坏，或是旧版 .doc 格式，请另存为 .docx）：{error}") from error
    style_names = {style.style_id: style.name or "" for style in document.styles}
    paragraphs = list(document.element.body.iter(qn("w:p")))
    units = []
    for position, paragraph in enumerate(paragraphs):
        if position % 50 == 0:
            progress(position + 1, len(paragraphs), "paragraph")
        number = len(units) + 1
        line, images = [], []
        # Walk body XML so text, tables and inline/floating images stay in document order.
        for node in paragraph.iter():
            if node.tag == qn("w:t") and node.text:
                line.append(node.text)
            elif node.tag in (qn("w:tab"), qn("w:br")):
                line.append("\t" if node.tag == qn("w:tab") else "\n")
            elif node.tag == qn("a:blip"):
                rel_id = node.get(qn("r:embed"))
                part = document.part.related_parts.get(rel_id) if rel_id else None
                if part is not None:
                    text = ocr.image_text(part.blob)
                    images.append((part.blob, text))
                    line.append("\n" + text + "\n")
        text = "".join(line).strip()
        if not text and not images:
            continue
        for blob, ocr_text in images:
            figures.save_blob(blob, number, ocr_text)
        level = _docx_heading_level(paragraph, qn, style_names) if text and "\n" not in text else None
        units.append(Unit(number, f"第 {number} 段", [Block(text or "（图片）", heading=level,
                                                            source="styles" if level else None)]))
        if hasattr(progress, "page"):
            progress.page(units[-1], [i for i in figures.items if i["page"] == number], len(paragraphs))
    source = chapters.detect(units, "docx")
    return units, source, "paragraph"


# --------------------------------------------------------------------------- PowerPoint

def _pptx_sections(presentation):
    """Returns {slide_number: section_name} from the PowerPoint 2010 section list."""
    try:
        root = presentation.part._element
        ids = [int(s.get("id")) for s in presentation.slides._sldIdLst]
        sections = {}
        for section in root.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}section"):
            name = (section.get("name") or "").strip()
            members = [int(s.get("id")) for s in section.iter(
                "{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId")]
            if name and members and members[0] in ids:
                sections[ids.index(members[0]) + 1] = name
        # PowerPoint adds a "Default Section"/"默认节" when sections are unused.
        if len(sections) == 1 and list(sections.values())[0].lower() in ("default section", "默认节"):
            return {}
        return sections
    except Exception:
        return {}


def parse_pptx(path, ocr, figures, progress):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        presentation = Presentation(path)
    except Exception as error:
        raise DocumentError(f"无法读取 PPT 文件（可能已损坏，或是旧版 .ppt 格式，请另存为 .pptx）：{error}") from error
    sections = _pptx_sections(presentation)

    def shape_blocks(shapes, number, skip_id):
        blocks = []
        for shape in shapes:
            if shape.shape_id == skip_id:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                blocks.extend(shape_blocks(shape.shapes, number, skip_id))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (
                    shape.is_placeholder and hasattr(shape, "image")):
                try:
                    blob = shape.image.blob
                except Exception:
                    continue
                text = ocr.image_text(blob)
                figures.save_blob(blob, number, text)
                if text.strip():
                    blocks.append(Block(text, kind="ocr"))
            elif shape.has_table:
                rows = ["\t".join(cell.text for cell in row.cells if not cell.is_spanned)
                        for row in shape.table.rows]
                blocks.append(Block("\n".join(rows), kind="table"))
            elif shape.has_text_frame and shape.text.strip():
                blocks.append(Block(shape.text.strip()))
        return blocks

    units = []
    slides = list(presentation.slides)
    for number, slide in enumerate(slides, 1):
        progress(number, len(slides), "slide")
        blocks = []
        if number in sections:
            blocks.append(Block(sections[number], heading=1, kind="section", source="slides"))
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        if title:
            blocks.append(Block(title, kind="title"))
        body = shape_blocks(slide.shapes, number, title_shape.shape_id if title_shape is not None else None)
        if not title and body and body[0].kind == "text" and len(body[0].text.splitlines()[0]) <= 40:
            # No title placeholder: the first short text line usually acts as the title.
            first = body[0].text.splitlines()
            blocks.append(Block(first[0], kind="title"))
            body[0] = Block("\n".join(first[1:]))
        blocks.extend(b for b in body if b.text.strip())
        units.append(Unit(number, f"第 {number} 张幻灯片", blocks,
                          meta={"layout": slide.slide_layout.name or ""}))
        if hasattr(progress, "page"):
            progress.page(units[-1], [i for i in figures.items if i["page"] == number], len(slides))
    source = chapters.detect(units, "pptx")
    return units, source, "slide"


# --------------------------------------------------------------------------- Text / Markdown

def decode_text(data):
    """UTF-8 first; otherwise detect the encoding (GBK/GB18030 and Big5 are common
    for Chinese course notes). charset-normalizer is MIT-licensed."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    from charset_normalizer import from_bytes
    best = from_bytes(data).best()
    if best is not None and best.encoding:
        return str(best)
    for encoding in ("gb18030", "big5"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("无法识别文本编码，请另存为 UTF-8 后导入")


MD_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
FENCE = re.compile(r"^\s{0,3}(```|~~~)")


def parse_text(path, markdown, progress, max_lines=40):
    lines = decode_text(Path(path).read_bytes()).splitlines()
    units = []
    current, start = [], 0

    def flush(end):
        nonlocal current
        if current and any(l.strip() for l, _ in current):
            heading = current[0][1] if len(current) == 1 else None
            text = "\n".join(l for l, _ in current).strip("\n")
            number = len(units) + 1
            first, last = start + 1, end
            label = f"第 {first} 行" if first == last else f"第 {first}–{last} 行"
            units.append(Unit(number, label, [Block(text, heading=heading,
                                                    source="markdown" if heading else None)],
                              loc_start=first, loc_end=last))
            if hasattr(progress, "page"):
                progress.page(units[-1], [], 0)
        current = []

    in_fence = False
    for index, raw_line in enumerate(lines):
        line = raw_line.rstrip()
        if markdown and FENCE.match(line):
            in_fence = not in_fence
        heading = None
        if markdown and not in_fence:
            match = MD_HEADING.match(line)
            if match:
                heading = len(match.group(1))
                # Chapter titles and the reader show the heading without its "#" marks.
                line = match.group(2).strip()
            elif (re.fullmatch(r"\s{0,3}(=+|-+)\s*", line) and current and len(current) == 1
                  and current[0][0].strip() and index > 0):
                # Setext heading: previous single line underlined with === or ---.
                current[0] = (current[0][0], 1 if line.strip()[0] == "=" else 2)
                flush(index + 1)
                start = index + 1
                continue
        if heading or (not in_fence and chapters.looks_like_heading(line)):
            # Headings become their own unit so chapters start exactly at that line.
            flush(index)
            start = index
            current = [(line, heading)]
            flush(index + 1)
            start = index + 1
            continue
        if not line.strip() and not in_fence:
            flush(index)
            start = index + 1
            continue
        if not current:
            start = index
        current.append((line, None))
        if len(current) >= max_lines and not in_fence:
            flush(index + 1)
            start = index + 1
    flush(len(lines))
    progress(1, 1, "text")
    source = chapters.detect(units, "md" if markdown else "txt")
    return units, source, "line"


# --------------------------------------------------------------------------- entry points

def plain_text(units, loc_kind):
    sections = []
    for unit in units:
        body = unit.text.strip()
        if not body:
            continue
        sections.append(f"[{unit.label}]\n{body}" if loc_kind in ("page", "slide") else body)
    return ("\n\n" if loc_kind in ("page", "slide", "line") else "\n").join(sections)


def parse_structured(path, figures_dir=None, progress=None, cache_dir=None, ocr_mode="auto"):
    path = Path(path)
    ocr = Ocr(ocr_mode, cache_dir)
    figures = Figures(figures_dir)
    progress = progress or Progress(False)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        units, source, loc_kind = parse_pdf(path, ocr, figures, progress, cache_dir)
    elif suffix == ".docx":
        units, source, loc_kind = parse_docx(path, ocr, figures, progress)
    elif suffix == ".pptx":
        units, source, loc_kind = parse_pptx(path, ocr, figures, progress)
    elif suffix in (".txt", ".md", ".markdown"):
        units, source, loc_kind = parse_text(path, suffix != ".txt", progress)
    else:
        raise DocumentError("不支持的文件格式")
    content = plain_text(units, loc_kind).strip()
    if not content:
        raise DocumentError("未提取到正文（已对图片尝试 OCR），请检查资料内容或扫描清晰度")
    warnings = []
    suspect = sum(bool(b.issues) for u in units for b in u.blocks)
    if suspect:
        warnings.append(f"{suspect} 处内容含可疑 Unicode 映射，已保留提取原文，请对照原页面核验公式")
    if figures.skipped:
        warnings.append(f"{figures.skipped} 张图片因格式不支持或数量过多未保存")
    if loc_kind == "page" and units and sum(u.ocr for u in units) == len(units):
        warnings.append("全部页面为扫描件，正文来自 OCR，可能存在识别错误")
    return {
        "content": content,
        "loc_kind": loc_kind,
        "structure_source": source,
        "pages": [unit.to_json() for unit in units],
        "images": figures.items,
        "warnings": warnings,
    }


def parse_document(path):
    """Plain text only (phase-one interface)."""
    return parse_structured(path)["content"]
