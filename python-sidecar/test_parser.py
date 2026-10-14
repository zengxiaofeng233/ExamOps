import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from PIL import Image, ImageDraw, ImageFont
from PyPDF2 import PdfReader, PdfWriter
from PyPDF2.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from parser import DocumentError, Ocr, decode_text, parse_document, parse_structured

CJK_FONTS = [
    'C:/Windows/Fonts/msyh.ttc',
    '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
    '/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc',
    '/System/Library/Fonts/PingFang.ttc',
]


def cjk_font(size):
    """A font that can draw Chinese test images on Windows, Linux or macOS."""
    for candidate in CJK_FONTS:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    raise unittest.SkipTest('未找到可绘制中文的字体（Windows 自带微软雅黑；Linux 请安装 fonts-noto-cjk）')


def make_pdf(path, image=None, native="Native title", pages=1):
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=600, height=800)
        page = writer.pages[-1]
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        resources = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        commands = f'BT /F1 22 Tf 40 740 Td ({native}) Tj ET\n'.encode('ascii') if native else b''
        if image is not None:
            stream = DecodedStreamObject()
            stream.set_data(image.convert('RGB').tobytes())
            stream = stream.flate_encode()
            stream.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                           NameObject('/Width'): NumberObject(image.width), NameObject('/Height'): NumberObject(image.height),
                           NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
            resources[NameObject('/XObject')] = DictionaryObject({NameObject('/Im1'): writer._add_object(stream)})
            commands += b'q 520 0 0 130 40 500 cm /Im1 Do Q\n'
        page[NameObject('/Resources')] = resources
        contents = DecodedStreamObject()
        contents.set_data(commands)
        page[NameObject('/Contents')] = writer._add_object(contents)
    with open(path, 'wb') as file:
        writer.write(file)


def make_fixtures(folder):
    image = Image.new('RGB', (1100, 230), 'white')
    draw = ImageDraw.Draw(image)
    font = cjk_font(42)
    for y, text in [(20, '简体中文课程资料'), (85, '繁體中文課程資料'), (150, 'English study material 123')]:
        draw.text((25, y), text, font=font, fill='black')
    image.save(folder / '图片.png')
    (folder / '中文资料.txt').write_text('简体中文\n繁體中文\nEnglish study', encoding='utf-8-sig')
    (folder / '笔记.md').write_text('# 学习笔记\n正文内容', encoding='utf-8')
    (folder / '旧编码.txt').write_bytes(('中文课程资料，考试复习。' * 30).encode('gb18030'))
    make_pdf(folder / '混合.pdf', image)
    make_pdf(folder / '扫描.pdf', image, native='', pages=2)
    make_pdf(folder / '文字.pdf')
    make_pdf(folder / '空白.pdf', native='')
    document = Document()
    document.add_paragraph('Word 原生正文')
    document.add_table(rows=1, cols=2).cell(0, 0).text = '表格内容'
    document.add_picture(str(folder / '图片.png'))
    document.save(folder / '讲义.docx')
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(6), Inches(1)).text = 'PPT 原生正文'
    slide.shapes.add_picture(str(folder / '图片.png'), Inches(0.2), Inches(1.2), width=Inches(6))
    presentation.save(folder / '课件.pptx')


class ParserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix='examops-parser-')
        cls.folder = Path(cls.directory.name)
        make_fixtures(cls.folder)

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def test_text_and_markdown_without_ocr(self):
        with patch.object(Ocr, 'lines', side_effect=AssertionError('Text must not start OCR')):
            self.assertIn('繁體中文', parse_document(self.folder / '中文资料.txt'))
            notes = parse_document(self.folder / '笔记.md')
            self.assertTrue(notes.startswith('学习笔记'), notes)  # heading text without "#"
            self.assertIn('正文内容', notes)
            self.assertIn('中文课程资料', parse_document(self.folder / '旧编码.txt'))
            self.assertIn('Native title', parse_document(self.folder / '文字.pdf'))

    def test_ocr_offline_mixed_pdf_deduplicates_native_text(self):
        with patch('socket.socket.connect', side_effect=AssertionError('OCR must be offline')):
            text = parse_document(self.folder / '混合.pdf')
        self.assertEqual(text.count('Native title'), 1)
        self.assertIn('简体中文', text)
        self.assertIn('繁體中文', text)
        self.assertIn('English study material 123', text)

    def test_scanned_pages_are_both_preserved(self):
        text = parse_document(self.folder / '扫描.pdf')
        self.assertIn('[第 2 页]', text)
        self.assertEqual(text.count('English study material 123'), 2)

    def test_office_text_tables_and_embedded_images(self):
        word = parse_document(self.folder / '讲义.docx')
        self.assertIn('Word 原生正文', word)
        self.assertIn('表格内容', word)
        self.assertIn('繁體中文', word)
        slides = parse_document(self.folder / '课件.pptx')
        self.assertIn('PPT 原生正文', slides)
        self.assertIn('English study material 123', slides)

    def test_empty_corrupt_unsupported_are_errors(self):
        with self.assertRaisesRegex(ValueError, '未提取到正文'):
            parse_document(self.folder / '空白.pdf')
        broken = self.folder / '损坏.pdf'
        broken.write_bytes(b'not a PDF')
        with self.assertRaises(Exception):
            parse_document(broken)
        with self.assertRaisesRegex(ValueError, '不支持'):
            parse_document(self.folder / '图片.png')

    def test_json_protocol_and_process_exit(self):
        request = json.dumps({'command': 'parse_document', 'data': {'path': str(self.folder / '中文资料.txt')}})
        result = subprocess.run([sys.executable, str(Path(__file__).with_name('main.py'))], input=request + '\n',
                                encoding='utf-8', capture_output=True, timeout=30, check=True)
        ready, body = map(json.loads, result.stdout.splitlines())
        self.assertEqual(ready['status'], 'ready')
        self.assertEqual(body['status'], 'success')
        self.assertIn('繁體中文', body['content'])



class EncodingTests(unittest.TestCase):
    def test_chinese_legacy_encodings_are_detected(self):
        text = '第一章 函数与极限\n函数是两个非空数集之间的对应关系，极限描述变化趋势。\n' * 3
        self.assertEqual(decode_text(text.encode('utf-8')), text)
        self.assertEqual(decode_text(('\ufeff' + text).encode('utf-8')), text)
        self.assertEqual(decode_text(text.encode('gb18030')), text)
        traditional = '第一章 函數與極限\n函數是兩個非空數集之間的對應關係，極限描述變化趨勢。\n' * 3
        self.assertEqual(decode_text(traditional.encode('big5')), traditional)

    def test_binary_garbage_is_rejected_or_decoded_without_crashing(self):
        try:
            decode_text(bytes(range(256)) * 4)
        except DocumentError as error:
            self.assertIn('编码', str(error))


class AbnormalFileTests(unittest.TestCase):
    """Files students really import: half-downloaded, renamed, restricted or oddly shaped.
    (Password-protected PDFs and the JSON error protocol are covered in test_structure.py.)"""

    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix='examops-abnormal-')
        folder = cls.folder = Path(cls.directory.name)
        make_pdf(folder / '原版.pdf', pages=2)
        restricted = PdfWriter()
        for page in PdfReader(str(folder / '原版.pdf')).pages:
            restricted.add_page(page)
        restricted.encrypt('', owner_password='owner')
        with open(folder / '仅限制编辑.pdf', 'wb') as file:
            restricted.write(file)
        whole = (folder / '原版.pdf').read_bytes()
        (folder / '下载中断.pdf').write_bytes(whole[:len(whole) // 2])
        (folder / '改了扩展名.docx').write_text('这其实是一个文本文件', encoding='utf-8')
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, 'w') as bundle:
            bundle.writestr('readme.txt', 'not a presentation')
        (folder / '不是课件.pptx').write_bytes(archive.getvalue())
        (folder / '只有空行.md').write_text(' \n\n\t\n', encoding='utf-8')
        (folder / '超长一行.txt').write_text('函数的极限' * 200_000, encoding='utf-8')

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def parse(self, name):
        return parse_structured(self.folder / name, str(self.folder / 'figures'))

    def test_editing_restrictions_alone_do_not_block_reading(self):
        # Owner password only (no password needed to open): imported normally.
        self.assertEqual(len(self.parse('仅限制编辑.pdf')['pages']), 2)

    def test_broken_or_renamed_files_fail_with_a_reason(self):
        for name, reason in [('下载中断.pdf', '已损坏'), ('改了扩展名.docx', '旧版 .doc'),
                             ('不是课件.pptx', '旧版 .ppt'), ('只有空行.md', '未提取到正文')]:
            with self.subTest(name=name), self.assertRaisesRegex(DocumentError, reason):
                self.parse(name)

    def test_a_huge_single_line_is_kept_whole(self):
        self.assertEqual(len(self.parse('超长一行.txt')['content']), 1_000_000)


if __name__ == '__main__':
    unittest.main()
