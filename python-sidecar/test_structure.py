"""Tests for structured extraction: chapters, locations, figures, previews, CLI."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from PIL import Image, ImageDraw
from PyPDF2 import PdfWriter
from PyPDF2.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject, ArrayObject, TextStringObject
from docx import Document
from pptx import Presentation
from pptx.util import Inches

import chapters
from parser import DocumentError, Ocr, parse_structured, merge_lines, same_text
from structure import Block, Unit, unicode_issues
from render import render_pages

MAIN = str(Path(__file__).with_name("main.py"))


def make_text_pdf(path, pages, outline=(), image=None, encrypt=None):
    """pages: list of [(text, font_size), ...] drawn top-down (ASCII, Helvetica)."""
    writer = PdfWriter()
    for lines in pages:
        writer.add_blank_page(width=600, height=800)
        page = writer.pages[-1]
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        resources = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        commands, y = b'', 760
        for text, size in lines:
            line_y = 30 if text.isdigit() else y  # page-number fixtures belong in the footer
            commands += f'BT /F1 {size} Tf 40 {line_y} Td ({text}) Tj ET\n'.encode('ascii')
            y -= size + 14
        if image is not None:
            stream = DecodedStreamObject()
            stream.set_data(image.convert('RGB').tobytes())
            stream = stream.flate_encode()
            stream.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                           NameObject('/Width'): NumberObject(image.width), NameObject('/Height'): NumberObject(image.height),
                           NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
            resources[NameObject('/XObject')] = DictionaryObject({NameObject('/Im1'): writer._add_object(stream)})
            commands += b'q 300 0 0 150 40 80 cm /Im1 Do Q\n'
        page[NameObject('/Resources')] = resources
        contents = DecodedStreamObject()
        contents.set_data(commands)
        page[NameObject('/Contents')] = writer._add_object(contents)
    parents = {}
    for level, title, page_index in outline:
        parents[level] = writer.add_outline_item(title, page_index, parent=parents.get(level - 1))
    if encrypt:
        writer.encrypt(encrypt)
    with open(path, 'wb') as file:
        writer.write(file)


def chart_image(seed=0):
    image = Image.new('RGB', (320, 200), 'white')
    draw = ImageDraw.Draw(image)
    for index in range(6):
        height = 30 + ((index * 37 + seed * 11) % 140)
        draw.rectangle([20 + index * 48, 190 - height, 50 + index * 48, 190], fill=(40 + index * 30, 90, 200 - seed * 20))
    return image


def headings(result):
    return [(b['heading'], b['text']) for p in result['pages'] for b in p['blocks'] if b['heading']]


class NoOcr(unittest.TestCase):
    """Most structure tests do not need OCR; patch it out for speed."""

    def setUp(self):
        self._lines = Ocr.lines
        Ocr.lines = lambda self, image: []

    def tearDown(self):
        Ocr.lines = self._lines


class ChineseNumberTests(unittest.TestCase):
    def test_numerals(self):
        self.assertEqual(chapters.chinese_number('十二'), 12)
        self.assertEqual(chapters.chinese_number('二十三'), 23)
        self.assertEqual(chapters.chinese_number('一百零五'), 105)
        self.assertEqual(chapters.chinese_number('３'), 3)
        self.assertEqual(chapters.chinese_number('iv'), 4)
        self.assertIsNone(chapters.chinese_number('abc'))


class FidelityTests(unittest.TestCase):
    def test_scanned_headings_accept_compact_and_fullwidth_numbering(self):
        units = [Unit(1, '1', [Block('第一章函数与极限', kind='ocr'),
                              Block('１．１ 映射', kind='ocr')]),
                 Unit(2, '2', [Block('第二章导数', kind='ocr')])]
        self.assertEqual(chapters.detect(units, 'pdf'), 'patterns')
        self.assertEqual([b.heading for u in units for b in u.blocks], [1, 2, 1])
        self.assertEqual(units[0].blocks[1].text, '１．１ 映射')

    def test_math_characters_are_not_deduplicated_or_normalized(self):
        for a, b in [('x+y', 'x−y'), ('x²', 'x2'), ('A', 'a'), ('ℝ', 'R'), ('x=y', 'xy')]:
            self.assertFalse(same_text(a, b))
        self.assertTrue(same_text('x + y', 'x+y'))
        self.assertFalse(same_text('\ue000', '\ue000'))
        math = '𝑥² + ∫₀¹ f(x) dx ≠ ∑ αᵢ ∈ ℝ'
        self.assertEqual(unicode_issues(math), [])
        self.assertEqual(unicode_issues('x\ue000\ufffd'), ['U+E000', 'U+FFFD'])

    def test_formula_and_suspect_blocks_survive_paragraph_merging(self):
        formula = 'a + b + c + d + e + f = 𝑥²'
        unit = Unit(1, '1', [Block('A sufficiently long introductory sentence'),
                            Block(formula), Block('following explanation without punctuation')])
        merge_lines(unit)
        self.assertEqual(len(unit.blocks), 3)
        data = unit.to_json()['blocks'][1]
        self.assertEqual(data['text'], formula)
        self.assertEqual(data['raw_text'], formula)
        self.assertTrue(data['formula'])

    def test_outline_matches_ocr_without_duplicate_synthetic_heading(self):
        units = [Unit(1, '1', [Block('Chapter 1 Limits', kind='ocr')]),
                 Unit(2, '2', [Block('Chapter 2 Derivatives', kind='ocr')])]
        chapters.detect(units, 'pdf', [(0, 'Chapter 1 Limits', 0), (0, 'Chapter 2 Derivatives', 1)])
        self.assertEqual([len(u.blocks) for u in units], [1, 1])
        self.assertTrue(all(u.blocks[0].heading == 1 for u in units))

    def test_bookmark_chapters_and_ocr_subsections_are_combined(self):
        units = [Unit(1, '1', [Block('第一章函数', kind='ocr'), Block('第一节映射', kind='ocr')])]
        self.assertEqual(chapters.detect(units, 'pdf', [(0, '第一章函数', 0)]), 'mixed')
        self.assertEqual([b.heading for b in units[0].blocks], [1, 2])

    def test_pattern_subsection_respects_bookmark_part_level(self):
        units = [Unit(1, '1', [Block('Chapter 1 Functions'), Block('1.1 Domain', kind='ocr')])]
        chapters.detect(units, 'pdf', [(0, 'Part A', 0), (1, 'Chapter 1 Functions', 0)])
        self.assertEqual([b.heading for b in units[0].blocks], [1, 2, 3])

    def test_a_single_real_section_keeps_its_level(self):
        units = [Unit(1, '1', [Block('1.1 映射', kind='ocr')])]
        self.assertEqual(chapters.detect(units, 'pdf'), 'patterns')
        self.assertEqual(units[0].blocks[0].heading, 2)

    def test_repeated_header_does_not_remove_the_printed_chapter_title(self):
        units = [Unit(n, str(n), [Block('第一章函数', kind='ocr', bbox=[20, 760, 200, 780]),
                 Block('第一章函数' if n == 1 else '正文内容', kind='ocr', bbox=[40, 600, 400, 640])],
                 width=600, height=800) for n in range(1, 6)]
        chapters.detect(units, 'pdf')
        self.assertEqual(units[0].blocks[0].text, '第一章函数')
        self.assertEqual(units[0].blocks[0].heading, 1)
        self.assertEqual(len(units[0].blocks), 1)

    def test_repeated_numeric_formula_fragment_is_not_a_page_number(self):
        units = [Unit(n, str(n), [Block('1', kind='ocr', bbox=[200, 30, 210, 40])],
                      width=600, height=800) for n in range(1, 6)]
        chapters.strip_furniture(units)
        self.assertEqual([u.text for u in units], ['1'] * 5)

    def test_ocr_contents_with_separate_page_number_column_are_not_chapters(self):
        toc = Unit(1, '目录', [])
        for i in range(1, 6):
            y = 700 - i * 30
            toc.blocks.extend([Block(f'第{i}章 标题', kind='ocr', bbox=[40, y, 300, y+15]),
                               Block(str(i*10), kind='ocr', bbox=[500, y, 530, y+15])])
        body = Unit(2, '正文', [Block('第1章 标题', kind='ocr'), Block('正文内容')])
        chapters.detect([toc, body], 'pdf')
        self.assertTrue(all(b.heading is None for b in toc.blocks))
        self.assertEqual(body.blocks[0].heading, 1)

    def test_repeated_subsection_names_in_different_chapters_are_kept(self):
        units = [Unit(1, '1', [Block('第一章 函数'), Block('第一节 概述')]),
                 Unit(2, '2', [Block('第二章 导数'), Block('第一节 概述')])]
        chapters.detect(units, 'pdf')
        self.assertEqual([b.heading for u in units for b in u.blocks], [1, 2, 1, 2])


class StructureTests(NoOcr):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix='examops-structure-')
        cls.folder = Path(cls.directory.name)

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def test_pdf_unicode_math_including_non_bmp_is_preserved(self):
        D, N, S = DictionaryObject, NameObject, TextStringObject
        writer = PdfWriter()
        writer.add_blank_page(600, 800)
        page = writer.pages[-1]
        text = '𝑥²+∫₀¹f(x)dx≠α∈ℝ'
        cmap = DecodedStreamObject()
        mapping = '\n'.join(f'<{i:04X}> <{c.encode("utf-16-be").hex()}>' for i, c in enumerate(text, 1))
        cmap.set_data(('begincmap\n/CMapType 2 def\n1 begincodespacerange\n<0000> <FFFF>\nendcodespacerange\n'
                       + str(len(text)) + ' beginbfchar\n' + mapping + '\nendbfchar\nendcmap').encode())
        cid = D({N('/Type'): N('/Font'), N('/Subtype'): N('/CIDFontType2'), N('/BaseFont'): N('/Arial'),
                 N('/CIDSystemInfo'): D({N('/Registry'): S('Adobe'), N('/Ordering'): S('Identity'),
                                         N('/Supplement'): NumberObject(0)}), N('/DW'): NumberObject(600)})
        font = D({N('/Type'): N('/Font'), N('/Subtype'): N('/Type0'), N('/BaseFont'): N('/Arial'),
                  N('/Encoding'): N('/Identity-H'), N('/DescendantFonts'): ArrayObject([writer._add_object(cid)]),
                  N('/ToUnicode'): writer._add_object(cmap)})
        page[N('/Resources')] = D({N('/Font'): D({N('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(('BT /F1 16 Tf 40 700 Td <' + ''.join(f'{i:04X}' for i in range(1, len(text)+1)) + '> Tj ET').encode())
        page[N('/Contents')] = writer._add_object(stream)
        path = self.folder / 'unicode-math.pdf'
        with path.open('wb') as file:
            writer.write(file)
        block = parse_structured(path)['pages'][0]['blocks'][0]
        self.assertEqual(block['text'], text)
        self.assertEqual(block['raw_text'], text)
        self.assertEqual(block['issues'], [])
        self.assertTrue(block['formula'])
        self.assertIsNotNone(block['bbox'])

    def test_pdf_outline_marks_headings_and_inserts_missing_titles(self):
        path = self.folder / 'outline.pdf'
        make_text_pdf(path, [
            [('Chapter 1 Limits', 20), ('A limit describes the value a function approaches.', 11)],
            [('More about limits and continuity.', 11)],
            [('Derivatives', 20), ('The derivative measures change.', 11)],
        ], outline=[(0, 'Chapter 1 Limits', 0), (1, 'One-sided limits', 1), (0, 'Chapter 2 Derivatives', 2)])
        result = parse_structured(path)
        self.assertEqual(result['structure_source'], 'outline')
        self.assertEqual(result['loc_kind'], 'page')
        found = headings(result)
        self.assertIn((1, 'Chapter 1 Limits'), found)
        self.assertIn((2, 'One-sided limits'), found)  # not printed on the page: inserted
        self.assertEqual(result['pages'][1]['blocks'][0]['text'], 'One-sided limits')
        self.assertEqual([p['number'] for p in result['pages']], [1, 2, 3])
        self.assertEqual(result['pages'][2]['label'], '第 3 页')

    def test_pdf_patterns_skip_toc_and_running_headers(self):
        path = self.folder / 'patterns.pdf'
        header = ('Calculus Lecture Notes', 9)
        make_text_pdf(path, [
            [('Contents', 16), ('Chapter 1 Functions ........ 2', 11), ('1.1 Domain ........ 2', 11),
             ('Chapter 2 Limits ........ 4', 11), ('2.1 Definition ........ 4', 11)],
            [header, ('Chapter 1 Functions', 16), ('A function maps inputs to outputs.', 11), ('2', 9)],
            [header, ('1.1 Domain', 13), ('The domain is the set of inputs.', 11), ('3', 9)],
            [header, ('Chapter 2 Limits', 16), ('Limits describe approach.', 11), ('4', 9)],
            [header, ('2.1 Definition', 13), ('Epsilon and delta.', 11), ('5', 9)],
        ])
        result = parse_structured(path)
        self.assertEqual(result['structure_source'], 'patterns')
        self.assertEqual(headings(result), [(1, 'Chapter 1 Functions'), (2, '1.1 Domain'),
                                            (1, 'Chapter 2 Limits'), (2, '2.1 Definition')])
        texts = [b['text'] for p in result['pages'] for b in p['blocks']]
        self.assertNotIn('Calculus Lecture Notes', texts)  # running header removed
        self.assertNotIn('3', texts)  # bare page number removed
        self.assertIn('Contents', texts)  # TOC page kept as text

    def test_pdf_layout_fallback_uses_font_size(self):
        path = self.folder / 'layout.pdf'
        make_text_pdf(path, [
            [('Introduction to Graphs', 24), ('Graphs consist of vertices and edges.', 11)],
            [('Plain body text continues here.', 11)],
            [('Shortest Paths', 24), ('Dijkstra relaxes edges greedily.', 11)],
        ])
        result = parse_structured(path)
        self.assertEqual(result['structure_source'], 'layout')
        self.assertEqual(headings(result), [(1, 'Introduction to Graphs'), (1, 'Shortest Paths')])

    def test_encrypted_and_corrupt_pdf_messages(self):
        path = self.folder / 'locked.pdf'
        make_text_pdf(path, [[('Secret', 12)]], encrypt='pw')
        with self.assertRaisesRegex(DocumentError, '加密'):
            parse_structured(path)
        broken = self.folder / 'broken.pdf'
        broken.write_bytes(b'%PDF-1.4 not really')
        with self.assertRaisesRegex(DocumentError, '无法打开 PDF'):
            parse_structured(broken)

    def test_docx_heading_styles_and_paragraph_locations(self):
        path = self.folder / 'styles.docx'
        document = Document()
        document.add_heading('第一章 函数与极限', level=1)
        document.add_paragraph('函数是两个集合之间的对应关系。')
        document.add_heading('1.1 映射', level=2)
        document.add_paragraph('映射是函数概念的推广。')
        document.add_heading('第二章 导数', level=1)
        document.add_paragraph('导数描述变化率。')
        document.save(path)
        result = parse_structured(path)
        self.assertEqual(result['structure_source'], 'styles')
        self.assertEqual(result['loc_kind'], 'paragraph')
        self.assertEqual(headings(result), [(1, '第一章 函数与极限'), (2, '1.1 映射'), (1, '第二章 导数')])
        self.assertEqual(result['pages'][1]['label'], '第 2 段')

    def test_docx_without_styles_uses_chinese_patterns(self):
        path = self.folder / 'plain.docx'
        document = Document()
        for text in ['课程讲义', '第一章 集合', '集合是确定对象的总体。', '第一节 集合的表示', '列举法与描述法。',
                     '第二章 映射', '映射的定义。', '“第三章介绍了很多内容”不是标题。']:
            document.add_paragraph(text)
        document.save(path)
        result = parse_structured(path)
        self.assertEqual(result['structure_source'], 'patterns')
        self.assertEqual(headings(result), [(1, '第一章 集合'), (2, '第一节 集合的表示'), (1, '第二章 映射')])

    def test_markdown_headings_ignore_code_blocks(self):
        path = self.folder / 'notes.md'
        path.write_text('# 第一部分\n\n正文一\n\n```python\n# not a heading\nx = 1\n```\n\n## 小节 A\n内容 A\n\n标题二\n===\n\n结尾\n',
                        encoding='utf-8')
        result = parse_structured(path)
        self.assertEqual(result['structure_source'], 'markdown')
        # Section titles drop the "#" marks of ATX headings.
        self.assertEqual(headings(result), [(1, '第一部分'), (2, '小节 A'), (1, '标题二')])
        self.assertEqual(result['loc_kind'], 'line')
        labels = [p['label'] for p in result['pages']]
        self.assertEqual(labels[0], '第 1 行')
        self.assertIn('第 5–8 行', labels)  # fenced code kept together with its line range

    def test_text_patterns_split_heading_from_paragraph(self):
        path = self.folder / 'book.txt'
        path.write_text('第一章 极限\n数列极限的定义如下。\n继续说明。\n\n第一节 数列\n数列是按顺序排列的数。\n\n第二章 导数\n导数是变化率。\n',
                        encoding='utf-8')
        result = parse_structured(path)
        self.assertEqual(result['structure_source'], 'patterns')
        self.assertEqual(headings(result), [(1, '第一章 极限'), (2, '第一节 数列'), (1, '第二章 导数')])
        first = result['pages'][0]
        self.assertEqual((first['loc_start'], first['loc_end']), (1, 1))
        self.assertEqual(result['pages'][1]['label'], '第 2–3 行')

    def test_pptx_titles_sections_and_continued_slides(self):
        path = self.folder / 'deck.pptx'
        presentation = Presentation()
        for title, body in [('第一章 概率基础', '本章导读'), ('随机事件', '事件的定义'), ('随机事件（续）', '事件的运算'),
                            ('第二章 随机变量', '本章导读'), ('分布函数', 'F(x)=P(X≤x)')]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        presentation.save(path)
        result = parse_structured(path)
        self.assertEqual(result['structure_source'], 'slides')
        self.assertEqual(result['loc_kind'], 'slide')
        self.assertEqual(headings(result), [(1, '第一章 概率基础'), (2, '随机事件'), (1, '第二章 随机变量'), (2, '分布函数')])
        self.assertEqual(result['pages'][4]['label'], '第 5 张幻灯片')

    def test_figures_are_saved_filtered_and_deduplicated(self):
        figures = self.folder / 'figures'
        pdf = self.folder / 'figure.pdf'
        make_text_pdf(pdf, [[('Chapter 1 Charts', 18)], [('Chapter 2 More', 18)]], image=chart_image())
        result = parse_structured(pdf, figures)
        # The same chart on both pages is kept once.
        self.assertEqual([image['page'] for image in result['images']], [1])
        saved = Path(result['images'][0]['path'])
        self.assertTrue(saved.exists() and saved.parent == figures)
        with Image.open(saved) as image:
            self.assertGreater(image.width, 100)

        image_path = self.folder / 'chart.png'
        chart_image(seed=2).save(image_path)
        Image.new('RGB', (20, 20), 'red').save(self.folder / 'tiny.png')
        document = Document()
        document.add_paragraph('图表说明')
        document.add_picture(str(image_path))
        document.add_picture(str(self.folder / 'tiny.png'))
        document.save(self.folder / 'figure.docx')
        result = parse_structured(self.folder / 'figure.docx', self.folder / 'docx-figures')
        self.assertEqual(len(result['images']), 1)  # tiny image filtered out
        self.assertEqual(result['images'][0]['page'], 2)

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(4), Inches(0.6)).text = '柱状图'
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(1), width=Inches(4))
        presentation.save(self.folder / 'figure.pptx')
        result = parse_structured(self.folder / 'figure.pptx', self.folder / 'pptx-figures')
        self.assertEqual([(i['page'], i['ordinal']) for i in result['images']], [(1, 1)])

    def test_render_pages_writes_previews(self):
        pdf = self.folder / 'render.pdf'
        make_text_pdf(pdf, [[('Page one', 14)], [('Page two', 14)]])
        pages = render_pages(pdf, [2, 9], self.folder / 'previews', width=400)
        self.assertEqual([p['page'] for p in pages], [2])
        self.assertEqual(pages[0]['width'], 400)
        self.assertTrue(Path(pages[0]['path']).exists())


class CliTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix='examops-cli-')
        cls.folder = Path(cls.directory.name)

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def run_cli(self, *args):
        return subprocess.run([sys.executable, MAIN, *args], encoding='utf-8', capture_output=True,
                              timeout=60, check=True)

    def test_input_returns_structure_and_reports_progress(self):
        pdf = self.folder / 'cli.pdf'
        make_text_pdf(pdf, [[('Chapter 1 Sets', 18), ('Body', 11)], [('Chapter 2 Maps', 18)]])
        result = self.run_cli('--input', str(pdf), '--progress')
        body = json.loads(result.stdout)
        self.assertEqual(body['status'], 'success')
        self.assertEqual(body['version'], 2)
        self.assertEqual(len(body['pages']), 2)
        progress = [json.loads(line)['progress'] for line in result.stderr.splitlines() if line.startswith('{"progress"')]
        self.assertEqual(progress[-1], {'stage': 'page', 'current': 2, 'total': 2})

    def test_errors_are_reported_as_json(self):
        missing = self.run_cli('--input', str(self.folder / 'missing.docx'))
        body = json.loads(missing.stdout)
        self.assertEqual(body['status'], 'error')
        self.assertIn('Word', body['message'])

    def test_render_cli(self):
        pdf = self.folder / 'cli-render.pdf'
        make_text_pdf(pdf, [[('Preview', 14)]])
        body = json.loads(self.run_cli('--render', str(pdf), '--pages', '1', '--out-dir',
                                       str(self.folder / 'out'), '--width', '300').stdout)
        self.assertEqual(body['status'], 'success')
        self.assertEqual(body['pages'][0]['width'], 300)


if __name__ == '__main__':
    unittest.main()
